import os
os.environ["HF_HUB_DISABLE_SYMLINKS_WARNING"] = "1"  # Disable symlink warnings for Hugging Face Hub

import streamlit as st
import pandas as pd
import docx
import PyPDF2
import nltk
from nltk.tokenize import sent_tokenize
import io
import re
import faiss
import pickle
import numpy as np
import sys
import zipfile
import shutil
from langchain_community.llms import Ollama
from langchain_ollama import ChatOllama
from langchain_core.prompts import PromptTemplate
from sentence_transformers import SentenceTransformer
from guardrails.hub import ToxicLanguage, NSFWText, GuardrailsPII
from guardrails import Guard
from fpdf import FPDF

nltk.download('punkt_tab')

st.set_page_config(page_title="Resume Screening", layout="wide")

# Dark mode CSS
theme_css = """
<style>
    .main {
        padding: 1rem;
    }
    .stApp {
        background: linear-gradient(135deg, #222831, #393E46);
        color: #ffffff;
        transition: background 0.3s, color 0.3s;
    }
    h1.title {
        text-align: center;
        color: #ffffff;
        font-size: 2.8rem;
        font-weight: bold;
        background: #948979;
        padding: 0.5rem;
        border-radius: 10px;
        margin-bottom: 1.5rem;
        text-shadow: 1px 1px 2px rgba(0, 0, 0, 0.2);
    }
    .stFileUploader {
        background-color: #393E46;
        border: 2px solid #948979;
        border-radius: 10px;
        padding: 1rem;
        margin-bottom: 1rem;
        transition: transform 0.2s, background-color 0.3s;
    }
    .stFileUploader:hover {
        transform: scale(1.02);
    }
    .stFileUploader label {
        color: #ffffff;
        font-size: 1.1rem;
        font-weight: 500;
    }
    .stFileUploader button {
        background-color: #948979 !important;
        color: #ffffff !important;
        border: 2px solid #DFD0B8 !important;
        padding: 0.5rem 1rem !important;
        border-radius: 8px !important;
        font-weight: 600 !important;
        font-size: 1rem !important;
    }
    .stFileUploader button:hover {
        background-color: #DFD0B8 !important;
        color: #222831 !important;
        transform: scale(1.05);
    }
    .stFileUploader section {
        border: 2px dashed #948979 !important;
        background-color: #2d3238 !important;
        border-radius: 8px !important;
    }
    .stFileUploader section:hover {
        border-color: #DFD0B8 !important;
        background-color: #353b42 !important;
    }
    .stFileUploader small {
        color: #DFD0B8 !important;
        font-size: 0.9rem !important;
    }
    .status-message {
        color: #ffffff;
        margin: 0.5rem 0;
        font-size: 1rem;
        padding: 0.5rem;
        background: #DFD0B8;
        border-radius: 5px;
    }
    .stButton button {
        background-color: #948979;
        color: #ffffff;
        border: none;
        padding: 0.7rem 1.5rem;
        border-radius: 20px;
        font-size: 1rem;
        transition: background-color 0.3s, transform 0.2s;
    }
    .stButton button:hover {
        background-color: #DFD0B8;
        transform: scale(1.05);
    }
    .download-container {
        background-color: #393E46;
        border: 2px solid #948979;
        border-radius: 10px;
        padding: 1rem;
        margin-top: 1.5rem;
    }
    .stMarkdown p, .stMarkdown div, .stMarkdown span, .stMarkdown h1, .stMarkdown h2, .stMarkdown h3, 
    .stMarkdown h4, .stMarkdown h5, .stMarkdown h6 {
        color: #ffffff;
    }
    .non-valid-count {
        color: #ffffff;
        background: #DFD0B8;
        padding: 0.5rem;
        border-radius: 5px;
        font-weight: bold;
        margin-top: 1rem;
    }
    .loading {
        display: inline-block;
        width: 20px;
        height: 20px;
        border: 3px solid #DFD0B8;
        border-top: 3px solid #948979;
        border-radius: 50%;
        animation: spin 1s linear infinite;
        margin-left: 10px;
    }
    @keyframes spin {
        to { transform: rotate(360deg); }
    }
</style>
"""
st.markdown(theme_css, unsafe_allow_html=True)

st.markdown("<h1 class='title'>Resume Screening</h1>", unsafe_allow_html=True)

@st.cache_resource
def load_model():
    return SentenceTransformer('all-MiniLM-L6-v2')

model = load_model()

# Initialize the Ollama LLM with Llama3.2 model
llm = ChatOllama(model="llama3.2")

column_names = [
    'SNO.', 'Full Name', 'Age', 'Email', 'Mobile Number',
    'Degree', 'Specialisation', 'Percentage', 'Total experience', 'skills', 'Applicant ID'
]

valid_specialisations = [
    "Artificial Intelligence", "Cloud Computing", "Computer", "Computer & Networking Security",
    "Computer And Information Science", "Computer Application", "Computer Science", "Data Science",
    "Electrical & Electronics", "Electronics", "Electronics & Communications",
    "Electronics & Telecom Engineering", "Electronics And Instrumentation",
    "Information Science And Engineering", "Information Technology", "Machine Learning",
    "Machine Learning And Data Science", "Software Engineering",
    "Computer Science And Engineering (cse)", "Electronics And Communications Engineering (ece)",
    "Cse", "Ece", "Cs"
]

def clean_resume_text(text):
    text = re.sub(r'[^\w\s\.\,\:\;\-\(\)\[\]/]', '', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def to_camel_case(text):
    if pd.isna(text):
        return ""
    return ' '.join(word.capitalize() for word in str(text).split())

def clean_degree_with_outliers(degree):
    if pd.isna(degree):
        return str(degree), False
    original = str(degree).strip()
    degree_lower = re.sub(r'[^a-z0-9]', '', original.lower())
    if re.search(r'baofte|bainte', degree_lower):
        return "B.Tech", True
    elif re.search(r'maofte|mainte', degree_lower):
        return "M.Tech", True
    elif re.search(r'maofcoap|maincoap', degree_lower):
        return "MCA", True
    elif re.search(r'baofsc|bainsc', degree_lower):
        return "B.Sc", True
    elif re.search(r'maofsc|mainsc', degree_lower):
        return "M.Sc", True
    elif re.search(r'doofph|doctora', degree_lower):
        return "PhD", True
    elif re.search(r'baofeng|baineng', degree_lower):
        return "B.E", True
    elif re.search(r'maofbuad|maofbu|mainbuad|mainbu', degree_lower):
        return "MBA", True
    if any(x in degree_lower for x in ["bacheloroftechnology", "bachelorstechnology", "bachelortechnology"]) or re.search(r'ba.*te', degree_lower):
        return "B.Tech", True
    if any(x in degree_lower for x in ["bachelorofeng", "bacheloreng", "bachelorsofeng"]):
        return "B.E", True
    if 'engineering' in degree_lower or 'engineer' in degree_lower:
        return "B.E/B.Tech", True
    if any(x in degree_lower for x in ["btech", "b.tech", "b-tech", "batchelor", "undergraduate", "bachelor'softechnology"]):
        return "B.Tech", True
    elif any(x in degree_lower for x in ["mtech", "masteroftechnology", "m.tech", "mastersoftechnology"]):
        return "M.Tech", True
    elif any(x in degree_lower for x in ["mca", "mastersofcomputerapplication", "masterofcomputerapplication", "masterIncomputerapplication", "mastersincomputerapplication"]):
        return "MCA", True
    elif any(x in degree_lower for x in ["bsc", "b.sc"]):
        return "B.Sc", True
    elif any(x in degree_lower for x in ["msc", "m.sc"]):
        return "M.Sc", True
    elif any(x in degree_lower for x in ["phd", "doctor", "doctoral"]):
        return "PhD", True
    elif any(x in degree_lower for x in ["diploma", "pgdiploma", "postgraduatediploma", "pgdbda", "dac", "cdac", "pgdca"]):
        return "Diploma", True
    elif any(x in degree_lower for x in ["be", "b.e", "bachelorofengineering", "bachelorsofengineering", "bachelor'sofengineering"]):
        return "B.E", True
    elif any(x in degree_lower for x in ["mba", "masterofbusiness"]):
        return "MBA", True
    elif any(x in degree_lower for x in ["ece", "cse"]):
        return "B.E/B.Tech", True
    elif any(x in degree_lower for x in ["bca", "bachelorofcomputerapplication", "bachelorincomputerapplication"]):
        return "BCA", True
    else:
        return original, False

def extract_degree_from_resume(resume_text, applicant_id):
    if not resume_text or not isinstance(resume_text, str):
        return None, False, None, None
    try:
        raw_resume = clean_resume_text(resume_text)
        prompt_template = PromptTemplate(
            input_variables=["resume_text"],
            template="""
            You are an assistant tasked with extracting only academic degree names from a resume. 
            Extract all academic degrees from the following list: B.Tech, M.Tech, M.Sc, B.Sc, MCA, BCA, B.E, M.E, Ph.D., Diploma.
            Return the degrees as a Python list (e.g., ["B.Tech", "M.Sc"]). 
            Do NOT include specializations (e.g., ignore 'CSE', 'Computer Science', 'ECE'). 
            Do NOT return Python code, regex patterns, or any text other than a Python list.
            For example:
            - Resume text contains 'BTech CSE' or 'B.Tech in Computer Science' -> return ["B.Tech"]
            - Resume text contains 'B.E - COMPUTER SCIENCE ENGINEERING' -> return ["B.E"]
            - Resume text contains 'M.Tech (ECE)' -> return ["M.Tech"]
            - No degrees found -> return []
            Incorrect outputs (do NOT return these):
            - Python code (e.g., def extract_degrees(...))
            - Tuples (e.g., [("B", "E")])
            - Text with specializations (e.g., ["BTech CSE"])
            Resume text: {resume_text}
            """
        )
        prompt = prompt_template.format(resume_text=raw_resume)
        response = llm.invoke(prompt)
        degrees = response.strip()
        valid_degrees = [
            'B.E', 'B.Tech', 'Bachelor of Science in Engineering',
            'Bachelor of Computer Science and Engineering', 'Bachelor of Information Technology',
            'Bachelor of Software Engineering', 'Bachelor of Electronics and Communication Engineering',
            'M.E', 'M.Tech', 'MS', 'Master of Software Engineering', 'Integrated M.Tech',
            'M.Sc', 'MCA', 'Ph.D.', 'B.E/B.Tech', 'Diploma', 'B.Sc', 'BCA'
        ]
        degree_list = []
        try:
            if degrees.startswith('[') and degrees.endswith(']'):
                degree_list = eval(degrees)
                if not isinstance(degree_list, list):
                    degree_list = []
            else:
                degree_list = []
        except Exception:
            degree_list = []
        if not degree_list:
            degree_pattern = r'\b(' + '|'.join([re.escape(deg) for deg in valid_degrees]) + r')\b'
            degree_list = re.findall(degree_pattern, raw_resume, re.IGNORECASE)
        cleaned_degree_list = []
        for degree in degree_list:
            for valid_degree in valid_degrees:
                if re.search(r'\b' + re.escape(valid_degree.lower()) + r'\b', degree.lower()):
                    cleaned_degree_list.append(valid_degree)
                    break
        degree_list = list(dict.fromkeys(cleaned_degree_list))
        toxic_validator = ToxicLanguage()
        nsfw_validator = NSFWText()
        metadata = {"applicant_id": applicant_id}
        validated_degrees = str(degree_list)
        try:
            toxic_result = toxic_validator.validate(validated_degrees, metadata)
            if isinstance(toxic_result, tuple):
                toxic_validation = toxic_result[1] if len(toxic_result) > 1 else True
            elif isinstance(toxic_result, dict):
                toxic_validation = toxic_result.get('is_valid', True)
            else:
                toxic_validation = True
            if not toxic_validation:
                return None, False, degree_list, "LLM response contains toxic or inappropriate language"
        except Exception:
            pass
        try:
            nsfw_result = nsfw_validator.validate(validated_degrees, metadata)
            if isinstance(nsfw_result, tuple):
                nsfw_validation = nsfw_result[1] if len(nsfw_result) > 1 else True
            elif isinstance(nsfw_result, dict):
                nsfw_validation = nsfw_result.get('is_valid', True)
            else:
                nsfw_validation = True
            if not nsfw_validation:
                return None, False, degree_list, "LLM response contains NSFW content"
        except Exception:
            pass
        if degree_list:
            if 'M.Tech' in degree_list:
                return 'M.Tech', True, degree_list, None
            for degree in degree_list:
                if degree in valid_degrees:
                    return degree, True, degree_list, None
        return None, False, degree_list, "No valid degrees found in resume"
    except Exception as e:
        st.error(f"Error extracting degree with LLM for Applicant ID {applicant_id}: {str(e)}")
        return None, False, None, f"Error in degree extraction: {str(e)}"

def experience_to_days(experience):
    if pd.isna(experience):
        return 0
    experience = str(experience).lower()
    years = months = days = 0
    try:
        years_match = re.search(r"(\d+)\s*year", experience, re.IGNORECASE)
        months_match = re.search(r"(\d+)\s*month", experience, re.IGNORECASE)
        days_match = re.search(r"(\d+)\s*day", experience, re.IGNORECASE)
        if years_match:
            years = int(years_match.group(1))
        if months_match:
            months = int(months_match.group(1))
        if days_match:
            days = int(days_match.group(1))
        return years * 365 + months * 30 + days
    except Exception as e:
        st.warning(f"Error parsing experience '{experience}': {str(e)}")
        return 0

def clean_skills(skills):
    if pd.isna(skills):
        return ""
    skills = str(skills)
    skill_list = []
    for part in skills.split(","):
        skill_name = str(part.strip().split("-")[0].strip())
        if skill_name:
            skill_list.append(skill_name)
    return ", ".join(skill_list)

def read_file(file):
    try:
        if isinstance(file, str):
            file_name = os.path.basename(file)
            if not os.path.exists(file):
                return None, f"File does not exist: {file_name}"
            if file_name.lower().endswith('.docx'):
                doc = docx.Document(file)
                text = ' '.join([para.text for para in doc.paragraphs if para.text.strip()])
                if text.strip():
                    return text.strip(), None
                return None, f"Empty document: {file_name}"
            elif file_name.lower().endswith('.pdf'):
                try:
                    with open(file, 'rb') as f:
                        pdf_reader = PyPDF2.PdfReader(f)
                        text = ""
                        for page in pdf_reader.pages:
                            page_text = page.extract_text() or ""
                            text += page_text + ' '
                        if text.strip():
                            return text.strip(), None
                        return None, f"No text extracted from {file_name}"
                except Exception as e:
                    if "encrypted" in str(e).lower():
                        return None, f"PDF is encrypted: {file_name}"
                    return None, f"PDF extraction failed for {file_name}: {str(e)}"
            elif file_name.lower().endswith('.txt'):
                with open(file, 'r', encoding='utf-8') as f:
                    text = f.read()
                    if text.strip():
                        return text.strip(), None
                    return None, f"Empty text file: {file_name}"
        else:  # Streamlit UploadedFile
            file_name = file.name
            if file_name.lower().endswith('.docx'):
                doc = docx.Document(file)
                text = ' '.join([para.text for para in doc.paragraphs if para.text.strip()])
                if text.strip():
                    return text.strip(), None
                return None, f"Empty document: {file_name}"
            elif file_name.lower().endswith('.pdf'):
                try:
                    pdf_reader = PyPDF2.PdfReader(file)
                    text = ""
                    for page in pdf_reader.pages:
                        page_text = page.extract_text() or ""
                        text += page_text + ' '
                    if text.strip():
                        return text.strip(), None
                    return None, f"No text extracted from {file_name}"
                except Exception as e:
                    if "encrypted" in str(e).lower():
                        return None, f"PDF is encrypted: {file_name}"
                    return None, f"PDF extraction failed for {file_name}: {str(e)}"
            elif file_name.lower().endswith('.txt'):
                text = file.read().decode('utf-8')
                if text.strip():
                    return text.strip(), None
                return None, f"Empty text file: {file_name}"
        return None, f"Unsupported file type: {file_name}"
    except Exception as e:
        return None, f"Error reading file {file_name}: {str(e)}"

def chunk_text(text, max_chunk_size=100, chunk_overlap=30):
    if not text or not isinstance(text, str):
        return []
    cleaned_text = re.sub(r'[^\w\s.,-]', ' ', text)
    cleaned_text = re.sub(r'\s+', ' ', cleaned_text).strip()
    try:
        sentences = sent_tokenize(cleaned_text)
        if not sentences and cleaned_text:
            return [cleaned_text]
        chunks = []
        current_chunk = ""
        for sentence in sentences:
            if len(current_chunk) + len(sentence) <= max_chunk_size:
                current_chunk += sentence + " "
            else:
                if current_chunk.strip():
                    chunks.append(current_chunk.strip())
                overlap_start = max(0, len(current_chunk) - chunk_overlap)
                current_chunk = current_chunk[overlap_start:] + sentence + " "
        if current_chunk.strip():
            chunks.append(current_chunk.strip())
        return chunks
    except Exception:
        if cleaned_text:
            return [cleaned_text]
        return []

def extract_info_from_pdf(pdf_text):
    try:
        info = {}
        full_text = " ".join(pdf_text.splitlines())
        full_text = re.sub(r'\s+', ' ', full_text)
        post_match = re.search(r'Post\s+([A-Za-z ()/]+)', full_text)
        if post_match:
            info['Post'] = post_match.group(1).strip()
        age_match = re.search(r'Age\s+(\d+)\s*years', full_text)
        if age_match:
            info['Age'] = age_match.group(1).strip() + " years"
        edu_match = re.search(r'1\)\s*(.*?)4\)\s*([^\)]+)(?=\s+[A-Z])', full_text)
        if edu_match:
            edu_1_to_3 = edu_match.group(1).strip()
            edu_4 = edu_match.group(2).strip()
            info['Educational Qualification'] = f"1) {edu_1_to_3} 4) {edu_4}"
        exp_match = re.search(r'Post\s+.*?relevant\s+Experience\s+([0-9]+(?:\s*years)?)', full_text, re.IGNORECASE)
        if exp_match:
            info['Post Qualification relevant Experience'] = exp_match.group(1).strip()
        return info
    except Exception as e:
        return {}

def get_llm_reason(candidate_row, requirements_text, screening_status, cleaned_valid_degrees):
    candidate_info = (
        f"Name: {candidate_row['Full Name']}\n"
        f"Age: {candidate_row['Age']}\n"
        f"Degree: {candidate_row['Degree']}\n"
        f"Specialisation: {candidate_row['Specialisation']}\n"
        f"Percentage: {candidate_row['Percentage']}\n"
        f"Experience: {candidate_row['Total experience']} days\n"
        f"Skills: {candidate_row['skills']}\n"
        f"Organization: {candidate_row.get('Organization', 'Unknown')}\n"
        f"Screening Status: {screening_status}"
    )
    try:
        candidate_degree = candidate_row['Degree'] if pd.notna(candidate_row['Degree']) else "Unknown"
        cleaned_candidate_degree = re.sub(r'[^a-z0-9]', '', str(candidate_degree).lower())
        organization = str(candidate_row.get('Organization', 'Unknown')).lower()
        is_cdac = bool(re.search(r'\bcdac\b', organization, re.IGNORECASE))
        total_experience = candidate_row['Total experience'] if pd.notna(candidate_row['Total experience']) else 0

        is_govt = False
        if organization != 'unknown':
            prompt_template = PromptTemplate(
                input_variables=["organization"],
                template="""
                You are an assistant tasked with determining if an organization is a government organization in India.
                Given the organization name: {organization}
                Return 'True' if the organization is a government entity (including CDAC), otherwise return 'False'.
                Examples:
                - CDAC -> True
                - ISRO -> True
                - DRDO -> True
                - Tata Consultancy Services -> False
                - Infosys -> False
                Do NOT return any explanation, only 'True' or 'False'.
                """
            )
            prompt = prompt_template.format(organization=organization)
            response = llm.invoke(prompt)
            is_govt = response.strip().lower() == 'true'

        if is_cdac or is_govt:
            if is_cdac and total_experience < 730:
                return f"Candidate was screened out because their organization is CDAC and total experience ({total_experience} days) is less than 2 years."
            else:
                if screening_status == "screened in":
                    return f"Candidate was screened in due to their valid degree, specialisation, sufficient percentage, and age. Candidate must submit a No Objection Certificate (NOC) at the time of interview as they are from a {'CDAC' if is_cdac else 'government'} organization."
                elif screening_status == "screened out":
                    reason = ""
                    if candidate_row['Specialisation'] not in valid_specialisations:
                        reason = f"Candidate was screened out because their specialisation '{candidate_row['Specialisation']}' is not in the required list."
                    elif pd.to_numeric(candidate_row['Age'], errors='coerce') > required_age:
                        reason = f"Candidate was screened out because their age ({candidate_row['Age']}) exceeds the required age limit ({required_age})."
                    elif pd.to_numeric(candidate_row['Percentage'], errors='coerce') < 60:
                        reason = f"Candidate was screened out because their percentage ({candidate_row['Percentage']}) is below the required 60%."
                    else:
                        reason = "Candidate was screened out for other reasons."
                    return f"{reason} Additionally, candidate must submit a No Objection Certificate (NOC) at the time of interview as they are from a {'CDAC' if is_cdac else 'government'} organization."
                elif screening_status == "manual check":
                    return f"Invalid degree or no valid degree found in resume. Candidate must submit a No Objection Certificate (NOC) at the time of interview as they are from a {'CDAC' if is_cdac else 'government'} organization."

        if screening_status == "manual check":
            return "Invalid degree or no valid degree found in resume"
        elif screening_status == "screened out" and candidate_row['Specialisation'] not in valid_specialisations:
            return f"Candidate was screened out because their specialisation '{candidate_row['Specialisation']}' is not in the required list."
        elif screening_status == "screened out" and pd.to_numeric(candidate_row['Age'], errors='coerce') > required_age:
            return f"Candidate was screened out because their age ({candidate_row['Age']}) exceeds the required age limit ({required_age})."
        elif screening_status == "screened out" and pd.to_numeric(candidate_row['Percentage'], errors='coerce') < 60:
            return f"Candidate was screened out because their percentage ({candidate_row['Percentage']}) is below the required 60%."
        elif screening_status == "screened in":
            return f"Candidate was screened in due to their valid degree, specialisation, sufficient percentage, and age."
        else:
            return "Screening reason could not be determined."
    except Exception as e:
        return f"Error generating reason: {str(e)}"

def generate_pdf_profile(profile_text, applicant_id, output_dir):
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, profile_text)
        pdf_file = os.path.join(output_dir, f"candidate_profile_{applicant_id}.pdf")
        pdf.output(pdf_file)
        return pdf_file
    except Exception as e:
        st.error(f"Error generating PDF for Applicant ID {applicant_id}: {str(e)}")
        return None

def generate_candidate_intro(profile_text, applicant_id):
    try:
        prompt_template = PromptTemplate(
            input_variables=["profile_text"],
            template="""
            You are an assistant tasked with generating a professional introduction for a candidate based on their profile.
            The introduction should be two paragraphs long, polished, and suitable for a hiring manager.
            Summarize the candidate's qualifications, experience, skills, and screening status concisely.
            Highlight their academic background, professional experience, and key strengths without including personally identifiable information (e.g., name, email, phone number).
            If the candidate is from a government organization or CDAC, mention the requirement for a No Objection Certificate (NOC) if applicable.
            Use a formal tone and avoid jargon or overly casual language.
            Example:
            - Input: "Degree: B.Tech\nSpecialisation: Computer Science\nPercentage: 75.0\nExperience: 1825 days\nSkills: Python, Java\nOrganization: Tech Corp\nScreening Status: screened in\nReason: Candidate was screened in due to their valid degree, specialisation, sufficient percentage, and age."
            - Output: This candidate holds a B.Tech degree in Computer Science with a commendable 75% academic score. With five years of professional experience, they have developed strong expertise in Python and Java, demonstrating proficiency in software development and problem-solving. Their qualifications align well with the job requirements, resulting in a successful screening outcome.

            The candidate's robust academic background and extensive experience make them a strong contender for the role. Their skills in Python and Java, combined with a proven track record in a reputable organization, position them as a valuable asset. They have been screened in due to their alignment with the required criteria, including degree, specialisation, and performance metrics.
            Profile text: {profile_text}
            """
        )
        prompt = prompt_template.format(profile_text=profile_text)
        response = llm.invoke(prompt)
        intro = response.strip()
        # Validate for toxic or NSFW content
        toxic_validator = ToxicLanguage()
        nsfw_validator = NSFWText()
        metadata = {"applicant_id": applicant_id}
        try:
            toxic_result = toxic_validator.validate(intro, metadata)
            toxic_validation = toxic_result.get('is_valid', True) if isinstance(toxic_result, dict) else True
            if not toxic_validation:
                return None, "Introduction contains toxic or inappropriate language"
        except Exception:
            pass
        try:
            nsfw_result = nsfw_validator.validate(intro, metadata)
            nsfw_validation = nsfw_result.get('is_valid', True) if isinstance(nsfw_result, dict) else True
            if not nsfw_validation:
                return None, "Introduction contains NSFW content"
        except Exception:
            pass
        return intro, None
    except Exception as e:
        return None, f"Error generating introduction: {str(e)}"

def save_candidate_intro(intro_text, applicant_id, output_dir):
    try:
        intro_file = os.path.join(output_dir, f"candidate_intro_{applicant_id}.txt")
        with open(intro_file, 'w', encoding='utf-8') as f:
            f.write(intro_text)
        return intro_file
    except Exception as e:
        st.error(f"Error saving introduction for Applicant ID {applicant_id}: {str(e)}")
        return None

def main():
    col1, col2 = st.columns(2)

    with col1:
        requirements_file = st.file_uploader("Requirements Document (.docx, .pdf, .txt)", type=['docx', 'pdf', 'txt'])
        excel_file = st.file_uploader("Resume Data (.xlsx)", type=['xlsx'])
        resume_folder = st.file_uploader("Resume Folder (.zip)", type=['zip'])

        if 'df_cleaned' not in st.session_state:
            st.session_state['df_cleaned'] = None
        if 'outlier_indices' not in st.session_state:
            st.session_state['outlier_indices'] = []
        if 'cleaned_rows' not in st.session_state:
            st.session_state['cleaned_rows'] = []
        if 'screened_data' not in st.session_state:
            st.session_state['screened_data'] = None
        if 'introductions' not in st.session_state:
            st.session_state['introductions'] = {}

        if excel_file and requirements_file and resume_folder:
            st.success("Documents uploaded successfully")
            try:
                full_df = pd.read_excel(excel_file, header=0)
                if full_df.empty:
                    st.error("Error: The Excel file is empty.")
                    raise ValueError("Empty Excel file")

                # Convert column names to lowercase and remove special characters
                full_df.columns = [re.sub(r'[^a-z0-9]', '', col.lower()) for col in full_df.columns]

                # Initialize column mappings
                column_mappings = {
                    'sno': None, 'fullname': None, 'age': None, 'email': None, 'mobilenumber': None,
                    'degree': None, 'specialisation': None, 'percentage': None, 'totalexperience': None,
                    'skills': None, 'otherexpertise': None, 'applicantid': None, 'organization': None
                }

                # Map columns with more flexible matching
                column_variants = {
                    'sno': ['sno', 'serialno', 'serialnumber', 'srno', 'slno', 'number', 'no'],
                    'fullname': ['fullname', 'name', 'candidatename', 'applicantname'],
                    'age': ['age'],
                    'email': ['email', 'emailid', 'emailaddress', 'mail'],
                    'mobilenumber': ['mobilenumber', 'mobile', 'phone', 'phonenumber', 'contact', 'contactnumber'],
                    'degree': ['degree', 'qualification', 'education'],
                    'specialisation': ['specialisation', 'specialization', 'branch', 'stream', 'major'],
                    'percentage': ['percentage', 'percent', 'marks', 'cgpa', 'gpa', 'score'],
                    'totalexperience': ['totalexperience', 'experience', 'exp', 'workexperience', 'yearsofexperience', 'experienceyears', 'totalexp'],
                    'skills': ['skills', 'skill', 'technicalskills', 'expertise'],
                    'otherexpertise': ['otherexpertise', 'additionalskills', 'otherskills', 'currentrole'],
                    'applicantid': ['applicantid', 'candidateid', 'applicationid', 'appid'],
                    'organization': ['organization', 'organisation', 'company', 'employer', 'currentcompany', 'govtcdac']
                }
                
                for col in full_df.columns:
                    col_clean = re.sub(r'[^a-z0-9]', '', col.lower())
                    for key, variants in column_variants.items():
                        if col_clean in variants and column_mappings[key] is None:
                            column_mappings[key] = col
                            break

                # If sno is still missing, fall back to applicantid column
                if column_mappings['sno'] is None and column_mappings['applicantid'] is not None:
                    column_mappings['sno'] = column_mappings['applicantid']

                # Check for missing required columns
                required_columns = ['sno', 'fullname', 'age', 'email', 'mobilenumber', 'degree',
                                   'specialisation', 'percentage', 'totalexperience', 'skills', 'applicantid']
                missing_columns = [col for col in required_columns if column_mappings[col] is None]
                if missing_columns:
                    st.error(f"Error: The following required columns are missing in the Excel file: {', '.join(missing_columns)}")
                    st.info(f"Columns found in your Excel file: {', '.join(full_df.columns.tolist())}")
                    st.warning("Please ensure your Excel file has columns for: Serial Number, Full Name, Age, Email, Mobile Number, Degree, Specialisation, Percentage, Total Experience, Skills, and Applicant ID")
                    st.session_state['df_cleaned'] = None
                    st.session_state['outlier_indices'] = []
                    st.session_state['cleaned_rows'] = []
                    st.session_state['screened_data'] = None
                    st.session_state['introductions'] = {}
                    raise ValueError("Missing required columns")

                # Rename columns safely
                column_name_map = {
                    'sno': 'SNO.', 'fullname': 'Full Name', 'age': 'Age', 'email': 'Email',
                    'mobilenumber': 'Mobile Number', 'degree': 'Degree', 'specialisation': 'Specialisation',
                    'percentage': 'Percentage', 'totalexperience': 'Total experience', 'skills': 'skills',
                    'applicantid': 'Applicant ID', 'organization': 'Organization'
                }

                # If sno and applicantid point to the same column, duplicate it first
                if column_mappings['sno'] == column_mappings['applicantid'] and column_mappings['sno'] is not None:
                    shared_col = column_mappings['sno']
                    full_df.insert(full_df.columns.get_loc(shared_col) + 1, 'SNO._temp', full_df[shared_col])
                    column_mappings['sno'] = 'SNO._temp'

                new_columns = full_df.columns.to_list()
                for key, value in column_mappings.items():
                    if value is not None and value in full_df.columns and key in column_name_map:
                        try:
                            idx = full_df.columns.get_loc(value)
                            new_columns[idx] = column_name_map[key]
                        except ValueError as e:
                            st.error(f"Error mapping column {value}: {str(e)}")
                            raise
                full_df.columns = new_columns

                # Check for SNO. column before duplicate check
                if 'SNO.' not in full_df.columns:
                    st.error("Error: 'SNO.' column not found after column mapping. Please ensure the Excel file contains a column like 'SNO', 'Serial No', or similar.")
                    raise KeyError("SNO. column missing after mapping")

                # Check for duplicate SNO. values
                if full_df['SNO.'].duplicated().any():
                    st.warning("Duplicate SNO. values found in Excel file. This may cause merge issues.")
                    st.write(f"Duplicate SNO. values: {full_df[full_df['SNO.'].duplicated()]['SNO.'].tolist()}")

                # Find the other expertise column
                skills_extra_col = column_mappings['otherexpertise']

                # Merge skills with other expertise
                if skills_extra_col is not None:
                    full_df['skills'] = full_df.apply(
                        lambda row: f"{row['skills']}, {row[skills_extra_col]}" if pd.notna(row[skills_extra_col]) else row['skills'],
                        axis=1
                    )
                    full_df = full_df.drop(columns=[skills_extra_col], errors='ignore')

                # Retain only relevant columns for cleaning
                df_cleaned = full_df[[col for col in column_names + ['Organization'] if col in full_df.columns]].copy()
                df_cleaned['LLM Degrees'] = ""
                df_cleaned['Screening Status'] = ""
                df_cleaned['Reason'] = ""

                valid_degrees = [
                    'B.E', 'B.Tech', 'Bachelor of Science in Engineering',
                    'Bachelor of Computer Science and Engineering', 'Bachelor of Information Technology',
                    'Bachelor of Software Engineering', 'Bachelor of Electronics and Communication Engineering',
                    'M.E', 'M.Tech', 'MS', 'Master of Software Engineering', 'Integrated M.Tech',
                    'M.Sc', 'MCA', 'Ph.D.', 'B.E/B.Tech', 'Diploma', 'B.Sc', 'BCA'
                ]
                cleaned_valid_degrees = [re.sub(r'[^a-z0-9]', '', deg.lower()) for deg in valid_degrees]

                cleaned_rows = []
                outlier_indices = []

                for index, row in df_cleaned.iterrows():
                    if pd.isna(row['Degree']) or pd.isna(row['Specialisation']):
                        st.warning(f"Row {index + 1} has missing Degree or Specialisation.")
                        cleaned_row = row.copy()
                        cleaned_row['Screening Status'] = "manual check"
                        cleaned_row['Reason'] = "Missing Degree or Specialisation"
                        cleaned_rows.append(cleaned_row)
                        continue
                    cleaned_row = row.copy()
                    std_deg, is_valid = clean_degree_with_outliers(cleaned_row['Degree'])
                    if is_valid:
                        cleaned_row['Degree'] = std_deg
                    if not is_valid:
                        outlier_indices.append(index + 1)
                        cleaned_row['Screening Status'] = "manual check"
                        cleaned_row['Reason'] = "Invalid degree"
                    cleaned_row['skills'] = clean_skills(cleaned_row['skills'])
                    cleaned_row['Total experience'] = experience_to_days(cleaned_row['Total experience'])

                    for col in df_cleaned.columns:
                        if isinstance(cleaned_row[col], str):
                            cleaned_row[col] = to_camel_case(cleaned_row[col])

                    cleaned_rows.append(cleaned_row)

                df_cleaned = pd.DataFrame(cleaned_rows, columns=df_cleaned.columns)

                # Check for duplicate SNO. in df_cleaned
                if 'SNO.' in df_cleaned.columns and df_cleaned['SNO.'].duplicated().any():
                    st.warning("Duplicate SNO. values found in cleaned DataFrame. This may cause merge issues.")
                    st.write(f"Duplicate SNO. values: {df_cleaned[df_cleaned['SNO.'].duplicated()]['SNO.'].tolist()}")

                # Process resume folder for degree extraction only for invalid or missing degrees
                temp_dir = "temp_resumes"
                try:
                    with zipfile.ZipFile(resume_folder, 'r') as zip_ref:
                        zip_ref.extractall(temp_dir)

                    for index, row in df_cleaned.iterrows():
                        applicant_id = str(row['Applicant ID']).strip()
                        candidate_degree = row['Degree'] if pd.notna(row['Degree']) else "Unknown"
                        cleaned_candidate_degree = re.sub(r'[^a-z0-9]', '', str(candidate_degree).lower())
                        
                        if cleaned_candidate_degree in cleaned_valid_degrees:
                            continue

                        resume_file = None
                        for root, dirs, files in os.walk(temp_dir):
                            for file_name in files:
                                norm_file_name = file_name.lower()
                                norm_applicant_id = applicant_id.lower()
                                if norm_applicant_id in norm_file_name and norm_file_name.endswith(('.pdf', '.docx', '.txt')):
                                    resume_file = os.path.join(root, file_name)
                                    break
                            if resume_file:
                                break
                        if not resume_file:
                            df_cleaned.at[index, 'Screening Status'] = "manual check"
                            df_cleaned.at[index, 'Reason'] = "Invalid degree or no resume found"
                            continue

                        resume_text, failure_reason = read_file(resume_file)
                        if resume_text and isinstance(resume_text, str) and resume_text.strip():
                            with st.spinner(f"Analyzing resume for Applicant ID {applicant_id} with Llama3.2..."):
                                new_degree, is_valid, degree_list, guardrail_reason = extract_degree_from_resume(resume_text, applicant_id)
                            if guardrail_reason:
                                df_cleaned.at[index, 'Screening Status'] = "manual check"
                                df_cleaned.at[index, 'Reason'] = guardrail_reason
                                if degree_list is not None:
                                    df_cleaned.at[index, 'LLM Degrees'] = str(degree_list)
                                continue
                            if degree_list is not None:
                                df_cleaned.at[index, 'LLM Degrees'] = str(degree_list)
                            if is_valid and new_degree:
                                df_cleaned.at[index, 'Degree'] = new_degree
                                df_cleaned.at[index, 'Screening Status'] = ""
                                df_cleaned.at[index, 'Reason'] = f"Degree updated from resume: {new_degree}"
                            else:
                                df_cleaned.at[index, 'Screening Status'] = "manual check"
                                df_cleaned.at[index, 'Reason'] = "No valid degree found in resume"
                        else:
                            df_cleaned.at[index, 'Screening Status'] = "manual check"
                            df_cleaned.at[index, 'Reason'] = f"Resume text extraction failed: {failure_reason}"

                except Exception as e:
                    st.error(f"Error processing resume folder: {str(e)}")
                    raise
                finally:
                    if os.path.exists(temp_dir):
                        for root, dirs, files in os.walk(temp_dir, topdown=False):
                            for file in files:
                                os.remove(os.path.join(root, file))
                            for dir in dirs:
                                os.rmdir(os.path.join(root, dir))
                        os.rmdir(temp_dir)
                    if os.path.exists("requirements_index.faiss"):
                        os.remove("requirements_index.faiss")
                    if os.path.exists("requirements_chunks.pkl"):
                        os.remove("requirements_chunks.pkl")

                try:
                    requirements_text = ""
                    extracted_info = {}
                    if requirements_file.name.endswith('.pdf'):
                        requirements_text, failure_reason = read_file(requirements_file)
                        if requirements_text:
                            extracted_info = extract_info_from_pdf(requirements_text)
                        else:
                            st.error(f"Failed to extract text from requirements file: {failure_reason}")
                            raise ValueError("Cannot process requirements file")
                    else:
                        requirements_text, failure_reason = read_file(requirements_file)
                        if not requirements_text:
                            st.error(f"Failed to extract text from requirements file: {failure_reason}")
                            raise ValueError("Cannot process requirements file")

                    if requirements_text and not os.path.exists("requirements_index.faiss"):
                        try:
                            chunks = chunk_text(requirements_text)
                            if chunks:
                                embeddings = model.encode(chunks, show_progress_bar=False)
                                if embeddings.size > 0:
                                    dimension = embeddings.shape[1]
                                    index = faiss.IndexFlatL2(dimension)
                                    index.add(embeddings)
                                    faiss.write_index(index, "requirements_index.faiss")
                                    with open("requirements_chunks.pkl", "wb") as f:
                                        pickle.dump(chunks, f)
                                else:
                                    st.error("No valid embeddings generated for requirements file.")
                                    raise ValueError("No valid embeddings")
                            else:
                                st.error("No valid chunks generated for requirements file.")
                                raise ValueError("No valid chunks")
                        except Exception as e:
                            st.error(f"Error processing requirements file: {str(e)}")
                            raise

                    is_fresher = False
                    if 'Post' in extracted_info and 'fresher' in extracted_info['Post'].lower():
                        is_fresher = True

                    global required_age
                    required_age = 100
                    if 'Age' in extracted_info:
                        age_match = re.search(r'(\d+)', extracted_info['Age'])
                        if age_match:
                            required_age = int(age_match.group(1))

                    st.write("Screening of the candidates going on 🌀")
                    for idx, row in df_cleaned.iterrows():
                        applicant_id = str(row['Applicant ID']).strip()
                        st.write(f"Processing Applicant ID {applicant_id} (Name: {row['Full Name']})")
                        candidate_degree = row['Degree'] if pd.notna(row['Degree']) else "Unknown"
                        cleaned_candidate_degree = re.sub(r'[^a-z0-9]', '', str(candidate_degree).lower())
                        if cleaned_candidate_degree not in cleaned_valid_degrees:
                            df_cleaned.at[idx, "Screening Status"] = "manual check"
                            df_cleaned.at[idx, "Reason"] = get_llm_reason(row, requirements_text, "manual check", cleaned_valid_degrees)
                            st.write(f"Applicant ID {applicant_id}: Set to manual check (invalid degree: {candidate_degree})")
                            continue

                        candidate_specialisation = row['Specialisation'] if pd.notna(row['Specialisation']) else "Unknown"
                        if candidate_specialisation not in valid_specialisations:
                            df_cleaned.at[idx, "Screening Status"] = "screened out"
                            df_cleaned.at[idx, "Reason"] = get_llm_reason(row, requirements_text, "screened out", cleaned_valid_degrees)
                            st.write(f"Applicant ID {applicant_id}: Screened out (invalid specialisation: {candidate_specialisation})")
                            continue

                        organization = str(row.get('Organization', 'Unknown')).lower()
                        is_cdac = bool(re.search(r'\bcdac\b', organization, re.IGNORECASE))
                        total_experience = row['Total experience'] if pd.notna(row['Total experience']) else 0
                        prompt_template = PromptTemplate(
                            input_variables=["organization"],
                            template="""
                            You are an assistant tasked with determining if an organization is a government organization in India.
                            Given the organization name: {organization}
                            Return 'True' if the organization is a government entity (including CDAC), otherwise return 'False'.
                            Examples:
                            - CDAC -> True
                            - ISRO -> True
                            - DRDO -> True
                            - Tata Consultancy Services -> False
                            - Infosys -> False
                            Do NOT return any explanation, only 'True' or 'False'.
                            """
                        )
                        prompt = prompt_template.format(organization=organization)
                        try:
                            response = llm.invoke(prompt)
                            is_govt = response.strip().lower() == 'true'
                        except Exception as e:
                            st.warning(f"Error checking organization for Applicant ID {applicant_id}: {str(e)}")
                            is_govt = False
                        if is_cdac and total_experience < 730:
                            df_cleaned.at[idx, "Screening Status"] = "screened out"
                            df_cleaned.at[idx, "Reason"] = get_llm_reason(row, requirements_text, "screened out", cleaned_valid_degrees)
                            st.write(f"Applicant ID {applicant_id}: Screened out (CDAC with experience {total_experience} days < 730)")
                            continue
                        elif is_govt or is_cdac:
                            if total_experience < 730:
                                df_cleaned.at[idx, "Screening Status"] = "screened out"
                                df_cleaned.at[idx, "Reason"] = get_llm_reason(row, requirements_text, "screened out", cleaned_valid_degrees)
                                st.write(f"Applicant ID {applicant_id}: Screened out (government/CDAC with experience {total_experience} days < 730)")
                                continue

                        if is_fresher:
                            age = pd.to_numeric(row['Age'], errors='coerce')
                            percentage = pd.to_numeric(row['Percentage'], errors='coerce')

                            if (pd.notna(age) and age <= required_age and
                                pd.notna(percentage) and percentage >= 60):
                                df_cleaned.at[idx, "Screening Status"] = "screened in"
                                df_cleaned.at[idx, "Reason"] = get_llm_reason(row, requirements_text, "screened in", cleaned_valid_degrees)
                                st.write(f"Applicant ID {applicant_id}: Screened in (fresher, age: {age}, percentage: {percentage})")
                            else:
                                df_cleaned.at[idx, "Screening Status"] = "screened out"
                                df_cleaned.at[idx, "Reason"] = get_llm_reason(row, requirements_text, "screened out", cleaned_valid_degrees)
                                st.write(f"Applicant ID {applicant_id}: Screened out (fresher, age: {age}, percentage: {percentage})")
                        else:
                            candidate_text = (
                                f"{row['Full Name']} {row['Degree']} {row['Specialisation']} "
                                f"{row['skills']} {row['Total experience']}"
                            )
                            try:
                                candidate_embedding = model.encode([candidate_text])[0]
                                candidate_embedding = candidate_embedding / np.linalg.norm(candidate_embedding)
                                candidate_embedding = candidate_embedding.reshape(1, -1).astype(np.float32)

                                if os.path.exists("requirements_index.faiss"):
                                    index = faiss.read_index("requirements_index.faiss")
                                    D, I = index.search(candidate_embedding, k=1)
                                    distance = D[0][0]
                                    cosine_similarity = 1 - (distance / 2)

                                    if cosine_similarity > 0.2:
                                        df_cleaned.at[idx, "Screening Status"] = "screened in"
                                        df_cleaned.at[idx, "Reason"] = get_llm_reason(row, requirements_text, "screened in", cleaned_valid_degrees)
                                        st.write(f"Applicant ID {applicant_id}: Screened in (cosine similarity: {cosine_similarity:.2f})")
                                    else:
                                        df_cleaned.at[idx, "Screening Status"] = "screened out"
                                        df_cleaned.at[idx, "Reason"] = get_llm_reason(row, requirements_text, "screened out", cleaned_valid_degrees)
                                        st.write(f"Applicant ID {applicant_id}: Screened out (cosine similarity: {cosine_similarity:.2f})")
                                else:
                                    df_cleaned.at[idx, "Screening Status"] = "screened out"
                                    df_cleaned.at[idx, "Reason"] = "Requirements not processed."
                                    st.write(f"Applicant ID {applicant_id}: Screened out (requirements not processed)")
                            except Exception as e:
                                df_cleaned.at[idx, "Screening Status"] = "manual check"
                                df_cleaned.at[idx, "Reason"] = f"Error in similarity check: {str(e)}"
                                st.write(f"Applicant ID {applicant_id}: Manual check (error in similarity check: {str(e)})")

                    # Merge cleaned data back into full_df to retain all original columns
                    full_df['Screening Status'] = ""
                    full_df['Reason'] = ""
                    full_df['LLM Degrees'] = ""

                    # Ensure merge by SNO. to update all rows
                    for idx, row in full_df.iterrows():
                        applicant_id = str(row['Applicant ID']).strip()
                        df_cleaned_row = df_cleaned[df_cleaned['SNO.'] == row['SNO.']]
                        if not df_cleaned_row.empty:
                            if len(df_cleaned_row) > 1:
                                st.warning(f"Multiple rows found for SNO. {row['SNO.']} in df_cleaned for Applicant ID {applicant_id}. Using first row.")
                            df_cleaned_row = df_cleaned_row.iloc[0]
                            for col in ['Degree', 'Specialisation', 'Percentage', 'Total experience', 'skills', 'Screening Status', 'Reason', 'LLM Degrees']:
                                if col in df_cleaned.columns and pd.notna(df_cleaned_row.loc[col]):
                                    full_df.at[idx, col] = df_cleaned_row.loc[col]
                            if 'Organization' in df_cleaned.columns and pd.notna(df_cleaned_row.get('Organization', '')):
                                full_df.at[idx, 'Organization'] = df_cleaned_row.loc['Organization']
                        else:
                            st.warning(f"Could not find matching SNO. {row['SNO.']} in df_cleaned for Applicant ID {applicant_id}")
                            full_df.at[idx, 'Screening Status'] = "manual check"
                            full_df.at[idx, 'Reason'] = "No matching row in cleaned data"

                    # Generate PDFs and introductions for each candidate with PII filtered
                    pdf_dir = "temp_pdfs"
                    intro_dir = "temp_introductions"
                    os.makedirs(pdf_dir, exist_ok=True)
                    os.makedirs(intro_dir, exist_ok=True)
                    pdf_files = []
                    intro_files = []
                    guard = Guard().use(GuardrailsPII(entities=["PHONE_NUMBER", "EMAIL_ADDRESS"], on_fail="fix"))
                    for idx, row in full_df.iterrows():
                        applicant_id = str(row['Applicant ID']).strip()
                        profile_text = (
                            f"Name: {row['Full Name']}\n"
                            f"Age: {row['Age']}\n"
                            f"Degree: {row['Degree']}\n"
                            f"Specialisation: {row['Specialisation']}\n"
                            f"Percentage: {row['Percentage']}\n"
                            f"Experience: {row['Total experience']} days\n"
                            f"Skills: {row['skills']}\n"
                            f"Organization: {row.get('Organization', 'Unknown')}\n"
                            f"Screening Status: {row['Screening Status']}\n"
                            f"Reason: {row['Reason']}"
                        )
                        try:
                            result = guard.validate(profile_text, metadata={"applicant_id": applicant_id})
                            filtered_profile = result.validated_output
                            pdf_file = generate_pdf_profile(filtered_profile, applicant_id, pdf_dir)
                            if pdf_file:
                                pdf_files.append(pdf_file)
                            # Generate and save introduction
                            intro_text, intro_error = generate_candidate_intro(filtered_profile, applicant_id)
                            if intro_text:
                                intro_file = save_candidate_intro(intro_text, applicant_id, intro_dir)
                                if intro_file:
                                    intro_files.append(intro_file)
                                    st.session_state['introductions'][applicant_id] = intro_text
                            else:
                                st.warning(f"Failed to generate introduction for Applicant ID {applicant_id}: {intro_error}")
                        except Exception as e:
                            st.warning(f"Failed to generate PDF or introduction for Applicant ID {applicant_id}: {str(e)}")

                    st.success("✅")

                    st.session_state['df_cleaned'] = full_df
                    st.session_state['outlier_indices'] = outlier_indices
                    st.session_state['cleaned_rows'] = cleaned_rows

                    if not full_df.empty:
                        zip_buffer = io.BytesIO()
                        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
                            excel_buffer = io.BytesIO()
                            with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
                                full_df.to_excel(writer, index=False, sheet_name='Screened Candidates')
                            zip_file.writestr('screened_candidates.xlsx', excel_buffer.getvalue())
                            for pdf_file in pdf_files:
                                zip_file.writestr(os.path.join('profiles', os.path.basename(pdf_file)), open(pdf_file, 'rb').read())
                            for intro_file in intro_files:
                                zip_file.writestr(os.path.join('introductions', os.path.basename(intro_file)), open(intro_file, 'rb').read())
                        st.session_state['screened_data'] = zip_buffer.getvalue()
                    else:
                        st.session_state['screened_data'] = None

                    # Clean up directories
                    if os.path.exists(pdf_dir):
                        shutil.rmtree(pdf_dir)
                    if os.path.exists(intro_dir):
                        shutil.rmtree(intro_dir)

                except Exception as e:
                    st.error(f"Error during screening at line {sys.exc_info()[-1].tb_lineno}: {str(e)}")
                    st.session_state['df_cleaned'] = None
                    st.session_state['outlier_indices'] = []
                    st.session_state['cleaned_rows'] = []
                    st.session_state['screened_data'] = None
                    st.session_state['introductions'] = {}
                    raise

            except Exception as e:
                st.error(f"Error processing Excel file at line {sys.exc_info()[-1].tb_lineno}: {str(e)}")
                st.session_state['df_cleaned'] = None
                st.session_state['outlier_indices'] = []
                st.session_state['cleaned_rows'] = []
                st.session_state['screened_data'] = None
                st.session_state['introductions'] = {}
                raise

    with col2:
        if requirements_file and excel_file and resume_folder:
            st.header("Download Results")
            if 'screened_data' in st.session_state and st.session_state['screened_data']:
                st.download_button(
                    label="Screened Candidates (Excel, Profiles, and Introductions)",
                    data=st.session_state['screened_data'],
                    file_name="screened_candidates.zip",
                    mime="application/zip"
                )
            else:
                st.write("No screened candidates available for download.")

            st.header("Candidate Introduction")
            candidate_name = st.text_input("Enter Candidate Name:")
            if candidate_name and 'df_cleaned' in st.session_state and st.session_state['df_cleaned'] is not None:
                full_df = st.session_state['df_cleaned']
                candidate_row = full_df[full_df['Full Name'].str.lower() == candidate_name.lower()]
                if not candidate_row.empty:
                    applicant_id = str(candidate_row.iloc[0]['Applicant ID']).strip()
                    if applicant_id in st.session_state['introductions']:
                        st.markdown("**Candidate Introduction:**")
                        st.markdown(st.session_state['introductions'][applicant_id])
                    else:
                        st.error(f"No introduction found for Applicant ID {applicant_id}")
                else:
                    st.error(f"No candidate found with name {candidate_name}")

if __name__ == "__main__":
    main()
